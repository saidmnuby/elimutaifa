from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).parent/'ElimuTaifa_Sponsors_Ads_Manual.docx'
d=Document()
s=d.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.7)
s.left_margin=s.right_margin=Inches(.8)
s.header_distance=s.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.12
d.styles['Title'].font.size=Pt(30)
d.styles['Heading 1'].font.size=Pt(18)
d.styles['Heading 2'].font.size=Pt(13)
for name in ['Heading 1','Heading 2']:
    d.styles[name].paragraph_format.space_before=Pt(12)
    d.styles[name].paragraph_format.space_after=Pt(8)
    d.styles[name].paragraph_format.keep_with_next=True
d.core_properties.title='ElimuTaifa Mwongozo wa Kuweka na Kusimamia Sponsors and Ads'
d.core_properties.subject='Mwongozo wa matumizi ya admin kwa direct sponsors na network ads'
d.core_properties.author='ElimuTaifa'
hp=s.header.paragraphs[0]; hp.text='ELIMUTAIFA  |  Sponsors & Ads'; hp.style='Normal'
hp.runs[0].font.size=Pt(9)
fp=s.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('Mwongozo wa admin  •  ')
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); fp._p.append(f)
for r in fp.runs: r.font.size=Pt(9)

def p(t): return d.add_paragraph(t)
def h(t): return d.add_heading(t,2)
def bullets(items):
    for t in items: d.add_paragraph(t,'List Bullet')
def steps(items):
    for i,t in enumerate(items,1): p(f'{i}. {t}')
def page(title):
    d.add_page_break(); d.add_heading(title,1)
def table(headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,txt in zip(t.rows[0].cells,headers): c.text=txt
    pr=t.rows[0]._tr.get_or_add_trPr(); flag=OxmlElement('w:tblHeader'); pr.append(flag)
    for row in rows:
        for c,txt in zip(t.add_row().cells,row): c.text=str(txt)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); no=OxmlElement('w:cantSplit'); trpr.append(no)
        for j,c in enumerate(row.cells):
            if widths: c.width=Inches(widths[j])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E9EEF2' if i==0 else ('F7F9FB' if i%2==0 else 'FFFFFF')); pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3); para.paragraph_format.space_before=Pt(3); para.paragraph_format.line_spacing=1.05
                for r in para.runs: r.font.size=Pt(10.5); r.bold=i==0
    p('')
    return t
def check(items):
    for t in items: p('☐  '+t)

# Cover
p('ELIMUTAIFA').runs[0].bold=True
d.add_paragraph('Mwongozo wa Kuweka na Kusimamia Sponsors and Ads','Title')
d.add_paragraph('Manual ya matumizi ya admin','Subtitle')
p('Toleo 1.0  |  4 Oktoba 2026')
h('Kwa main owner na wasimamizi wa matangazo')
p('Mwongozo huu unaeleza jinsi ya kuandaa maudhui ya sponsor, kuchagua kurasa, kupanga muda wa kuonekana na kukagua delivery. Unatumia majina ya fields na buttons ya English kama yanavyoonekana kwenye ElimuTaifa.')
p('Anza kwa kutambua aina ya tangazo. Direct sponsors na taarifa za mfumo huwekwa kupitia Direct placement. Google AdSense na Adsterra Native Banner husanidiwa kupitia Network ads & monitoring.')
h('Hali ya mfumo inayoelezwa')
p('Mwongozo unaeleza uwezo wa setup iliyopo. Hauonyeshi kuwa akaunti halisi za providers zimeunganishwa, website imeidhinishwa au matangazo yamepita connection test. Reporting APIs za impressions, clicks na earnings bado hazijaunganishwa.')
p('Hakuna screenshots, Publisher IDs, ad keys, mapato au approval zilizobuniwa kwenye document hii. Mfano wowote wa kujaza ni wa maelezo pekee.')

page('Yaliyomo')
contents=[('1','Madhumuni na aina za matangazo','3'),('2','URL na kurasa za matangazo','3'),('3','Kuweka direct sponsor','4–5'),('4','Platform announcements na Pop up','6'),('5','Kurasa priority status na schedule','7'),('6','Kuweka Google AdSense','8'),('7','Kuweka Adsterra Native Banner','9'),('8','Approval declaration na ukaguzi','10'),('9','Monitoring na kusimamia mabadiliko','11'),('10','Kutatatua matatizo','12'),('11','Checklist kabla ya kuchapisha','13'),('12','Daftari la sponsors','14')]
table(['Sehemu','Maudhui','Ukurasa'],contents,[.65,5.25,1.0])
p('Namba hizi zinafuata mpangilio wa toleo hili. Ukibadilisha maudhui au font kwenye Word, hakiki namba za yaliyomo kabla ya kuchapisha.')
h('Jinsi ya kutumia mwongozo')
bullets(['Direct sponsor: soma sehemu 3 hadi 5, kisha ukaguzi na checklist.','Google AdSense: soma sehemu 2, 5, 6 na 8 hadi 11.','Adsterra Native Banner: soma sehemu 2, 5, 7 na 8 hadi 11.','Tangazo likikosa kuonekana: tumia sehemu ya kutatua matatizo.'])
p('Njia kuu ya kufungua: Admin → Sponsors & Ads. Network units zinaweza kusanidiwa na main owner pekee; admin wengine wanaweza kusoma hali ya monitoring.')

page('1 Madhumuni na aina za matangazo')
p('Sponsors & Ads hutumika kuweka, kupanga, kuchapisha, kusimamisha na kukagua matangazo katika ElimuTaifa.')
table(['Aina','Sehemu ya kutumia','Vitu vya kuandaa'],[
('Direct sponsor','+ Direct placement','Jina, maelezo, picha na destination link'),('Platform announcement','+ Direct placement','Kichwa, maelezo na link au internal article'),('Google AdSense','Network ads & monitoring','Publisher ID na Ad slot'),('Adsterra Native Banner','Network ads & monitoring','Native Banner key na script URL')],[1.4,2.1,3.4])
p('Usiweke code ya Google au Adsterra kwenye Description, Image URL au External HTTPS destination.')
d.add_heading('2 URL na kurasa za matangazo',1)
p('Kwa usajili wa website kwa ad network, tumia URL kuu ya public website. Mfano wa maelezo pekee ni https://your-domain.com. Usitumie URL ya admin au localhost kama website ya public.')
p('Kwa kawaida huhitaji kusajili kila ukurasa mmoja mmoja. Katika ElimuTaifa, chagua kurasa kupitia Where to show au Display pages.')
table(['URL','Maana'],[('Public website URL','Website ya ElimuTaifa inayofunguliwa na wananchi'),('Destination URL','Sehemu mtumiaji anayoenda akibofya direct sponsor'),('Image HTTPS URL','Anwani ya picha ya tangazo'),('Native Banner script URL','Anwani ya script inayotolewa na Adsterra')],[2.05,4.85])
p('Nafasi na renderer ya matangazo tayari zipo. Huhitaji kuhariri code ya kila ukurasa unapoweka unit kupitia admin. Provider verification au consent inaweza kuhitaji setup nyingine kwenye production website.')

page('3 Kuweka direct sponsor')
steps(['Ingia admin na ufungue Sponsors & Ads.','Bofya + Direct placement.','Jaza taarifa za tangazo na uchague Type → Sponsor.','Chagua kurasa, nafasi na schedule.','Anza na Status → Draft, kisha bofya Save placement.','Tumia Preview kukagua maudhui.','Ukiridhika, badilisha kuwa Published, hifadhi na ukague public page.'])
h('Content and appearance')
table(['Field','Jinsi ya kujaza'],[
('Title','Kichwa cha tangazo, herufi 3–140'),('Description','Maelezo mafupi, hadi herufi 300'),('Type','Chagua Sponsor'),('Sponsor name','Jina la sponsor; ni lazima kwa Sponsor, hadi herufi 120'),('Format','Banner au Card'),('Position','Above content, Below content au Pop-up'),('Image HTTPS URL','Link ya picha inayotumia HTTPS'),('Or upload an image','JPEG, PNG au WebP, hadi 5 MB')],[2.0,4.9])
p('Kwa urahisi, tumia upload ya picha au Image HTTPS URL kama njia moja ya kuandaa picha. Maandishi yawe mafupi na yaeleze huduma au ofa bila kupotosha.')

page('3 Direct sponsor links na publishing')
h('Link and display pages')
table(['Field','Jinsi ya kujaza'],[
('Link type','External HTTPS link au Internal article'),('External HTTPS destination','Link halisi ya sponsor inayotumia HTTPS'),('Internal article','Chagua article ya ndani badala ya external link'),('Where to show','Chagua angalau ukurasa mmoja au kundi la kurasa')],[2.1,4.8])
h('Publishing and schedule')
table(['Field','Jinsi ya kujaza'],[('Status','Draft, Published au Archived'),('Priority','0–100; namba kubwa hupewa nafasi kwanza'),('Start / End','Tarehe na muda kwa EAT; si lazima. End iwe baada ya Start')],[2.1,4.8])
h('Mfano wa kujaza')
bullets(['Title: Jiunge na mafunzo ya kompyuta','Description: Kozi za msingi kwa wanafunzi na wahitimu.','Type: Sponsor; Sponsor name: jina halisi la kituo husika','Format: Banner; Position: Below content','Destination: HTTPS link halisi ya kituo','Where to show: Home na search pages','Priority: 50; Status: Draft wakati wa ukaguzi'])
p('Huu ni mfano wa maelezo. Hauwakilishi sponsor, link au tangazo lililowekwa tayari.')

page('4 Platform announcements na Pop up')
h('Kuweka platform announcement')
p('Fungua + Direct placement, kisha chagua Type → Platform announcement. Jaza kichwa na maelezo. Chagua Internal article ikiwa msomaji anapaswa kufungua taarifa ndani ya ElimuTaifa. Unaweza pia kutumia external HTTPS link inayofaa.')
p('Internal article lazima iwe published, imefikia publication time na haija-expire ili placement yake ionekane kwa public.')
h('Kutumia Pop up')
p('Pop-up inapatikana kwa direct placements na platform announcements. Network ad units katika setup hii zina nafasi za juu au chini pekee.')
table(['Field au option','Maana'],[
('Small corner pop-up','Tangazo dogo kwenye kona'),('Centre-screen announcement','Tangazo katikati ya screen'),('Once per session','Mara moja katika session'),('Every page visit','Kila page visit'),('Three times per browser','Hadi mara tatu kwa browser, kulingana na kumbukumbu yake'),('Wait before Skip is available','Sekunde 0–30 kabla ya Skip kupatikana')],[2.75,4.15])
p('Frequency inaweza kuanza upya mtumiaji akifuta browser storage au akitumia browser nyingine.')
p('Tumia pop-up kwa kiasi ili mtumiaji aendelee kupata matokeo au taarifa kwa urahisi. Hakiki style, frequency na Skip delay kabla ya kuchapisha.')

page('5 Kurasa priority status na schedule')
h('Kuchagua kurasa')
bullets(['All pages: kurasa zote za public zinazosaidia placements; haijumuishi admin pages.','Search pages, Result pages, School lists, Information pages au Error pages.','Kurasa maalumu kama Home, Announcements, ACSEE, CSEE, FTNA, PSLE, SFNA, Form One au Form Five, pamoja na kurasa zake zilizoorodheshwa kwenye form.'])
h('Nafasi na priority')
table(['Nafasi','Kikomo kwa ukurasa'],[('Above content','Placement moja'),('Below content','Placements mbili'),('Pop-up','Placement moja; direct placements pekee')],[2.0,4.9])
p('Direct sponsors na network ads hushirikiana katika nafasi za juu na chini. Priority kubwa hupewa nafasi kwanza. Tangazo linaweza kuwa Published lakini lisichaguliwe ikiwa nafasi zimechukuliwa na matangazo yenye priority kubwa.')
h('Status')
table(['Status','Tabia'],[('Draft','Imehifadhiwa lakini haionekani kwa public'),('Published','Huonekana ikiwa schedule, targets na nafasi vinaruhusu'),('Archived','Husimamisha kuonekana katika page loads zinazofuata')],[1.4,5.5])
h('Schedule kwa EAT')
bullets(['Start ikiwa tupu: hakuna muda wa kuanza unaozuia tangazo.','End ikiwa tupu: hakuna tarehe ya kumaliza iliyowekwa.','End lazima iwe baada ya Start.','Muda wa form ni EAT, saa za Tanzania.'])

page('6 Kuweka Google AdSense')
p('Main owner pekee anaweza kusanidi network units. Setup hii inasaidia responsive display ad units.')
steps(['Andaa public website na akaunti ya provider.','Kamilisha verification, approval na privacy/consent setup inayotakiwa na provider.','Pata responsive display ad code kwenye dashboard ya AdSense.','Fungua Sponsors & Ads → Network ads & monitoring, kisha form ya kuongeza ad unit.','Chagua Provider → Google AdSense.','Nakili Publisher ID na Ad slot kutoka kwenye code yako.'])
table(['Field ya ElimuTaifa','Taarifa ya kunakili'],[('Publisher ID','Thamani ya data-ad-client, inayoanza na ca-pub-'),('Ad slot / Native Banner key','Thamani ya data-ad-slot')],[2.8,4.1])
steps(['Weka Name inayotambulisha unit, mfano “AdSense — Search bottom”.','Chagua Position, Priority na Display pages.','Weka Start na End ikiwa zinahitajika.','Hifadhi Draft wakati bado unaandaa setup.','Baada ya approval na consent setup kukamilika, kagua declaration na uchapishe Published.'])
p('Mfumo hautaki kubandika script nzima kwenye form hii. AdSense units zilizochapishwa katika setup hii lazima zitumie Publisher ID moja.')
p('Kuhifadhi taarifa hakumaanishi Google imeidhinisha website au tangazo litaonekana mara moja. Hakuna Publisher ID halisi iliyowekwa kwenye mwongozo huu.')

page('7 Kuweka Adsterra Native Banner')
p('Setup hii inasaidia Native Banner. Haitumiki kubandika arbitrary scripts, Popunder au formats nyingine za Adsterra.')
steps(['Andaa website na ad unit kwenye dashboard ya Adsterra.','Pata Native Banner code.','Fungua Sponsors & Ads → Network ads & monitoring.','Chagua Provider → Adsterra Native Banner.','Jaza Name, Native Banner key na script URL halisi kutoka kwenye code ya provider.'])
table(['Field','Taarifa ya kuweka'],[('Name','Jina la kutambua unit'),('Ad slot / Native Banner key','Key ya herufi/namba 32 iliyo kwenye container-KEY'),('Native Banner script URL','HTTPS URL kutoka kwa provider inayoishia /KEY/invoke.js')],[2.7,4.2])
p('KEY hapa inaonyesha nafasi ya key yako halisi. Key kwenye script URL lazima ilingane na key iliyowekwa kwenye field ya unit.')
steps(['Chagua Position → Above content au Below content.','Weka Priority na Display pages.','Weka schedule ikiwa inahitajika.','Hifadhi Draft, au Published ikiwa approval na consent setup zimekamilika.'])
p('Usibuni key au script URL. Nakili taarifa halisi zinazotolewa kwenye dashboard yako. Application haithibitishi yenyewe umiliki wa URL kwa Adsterra.')

page('8 Approval declaration na ukaguzi')
h('Checkbox ya mwisho')
p('“My site and ad unit are approved, and provider-required privacy/consent setup is complete.”')
p('Main owner anathibitisha kwamba website na ad unit zina approval inayohitajika, na privacy/consent setup inayotakiwa na provider imekamilika.')
bullets(['Checkbox si connection test wala provider verification.','Haisakinishi consent-management tool.','Haihakikishi ad fill, kuonekana kwa tangazo au impression inayolipwa.','Usiweke tiki ikiwa hatua hizo bado hazijakamilika. Network unit haiwezi kuchapishwa bila declaration hii.'])
h('Kukagua direct sponsor')
steps(['Tumia Preview kukagua maandishi, picha na link.','Fungua public page uliyochagua.','Hakikisha tangazo linaonekana kwenye nafasi iliyokusudiwa.','Kagua destination link.','Kagua mwonekano kwenye simu na desktop.'])
p('Preview pekee haithibitishi kwamba schedule na targeting vinalifanya tangazo lionekane kwenye kila public page.')
h('Kukagua network ads')
bullets(['Localhost haipakii live network ads.','Kagua live delivery kwenye production website yenye setup halisi.','Kagua pia dashboard ya provider.','Script kupakiwa hakumaanishi tangazo limejazwa au impression imelipwa.','Usibofye matangazo yako kwa lengo la kufanya test.'])

page('9 Monitoring na kusimamia mabadiliko')
h('Maana ya local signals')
table(['Kipimo','Maana'],[('Mounted','Mfumo umeweka eneo la ad unit kwenye ukurasa'),('Script ready / Loaded','Script imepakiwa; si uthibitisho wa ad fill'),('Failed','Browser imeripoti tatizo la kupakia script')],[1.9,5.0])
p('Vipimo hivi ni makadirio ya afya ya delivery. Si takwimu rasmi za impressions, clicks au mapato. Marudio hupunguzwa kwa kila event katika kipindi cha saa moja kwa utambulisho wa muda wa browser/network. Bots na Do Not Track huondolewa kwenye local signals.')
p('Dashboard ina muhtasari wa siku saba za karibuni. Historia ya local signals huhifadhiwa hadi siku 60. Kutokuwepo kwa signal hakuthibitishi moja kwa moja kwamba provider ana hitilafu.')
p('Ripoti ya mfumo ya Ijumaa inaweza kujumuisha idadi ya network units na local script failures ikiwa weekly notifications zimewashwa.')
p('Reporting APIs za Google AdSense na Adsterra bado hazijaunganishwa. Takwimu rasmi za impressions, clicks na earnings zionekane kwenye dashboard ya provider.')
h('Kusimamisha tangazo')
steps(['Fungua tangazo kupitia Edit.','Badilisha Status kuwa Archived.','Hifadhi na ukague public page baada ya kuipakia tena.'])
p('Mabadiliko huathiri page loads zinazofuata. Tangazo lililopakiwa tayari kwenye browser iliyo wazi linaweza kubaki mpaka ukurasa upakiwe tena.')
h('Kubadilisha tangazo')
p('Kwa direct sponsor unaweza kubadilisha maudhui, picha, destination, targets, priority na schedule. Kwa network unit badilisha provider fields zinazofaa, targets, priority au schedule. Kagua tena baada ya kuhifadhi.')

page('10 Kutatua matatizo')
table(['Tatizo','Vitu vya kukagua'],[
('Tangazo halionekani','Status, targets, Start, End na priority'),('Linaonekana kwenye kurasa zisizotarajiwa','Je, All pages au kundi pana limechaguliwa?'),('Internal article placement haionekani','Article iwe published, imefikia publication time na haija-expire'),('Picha haionekani','File format, ukubwa hadi 5 MB au HTTPS image URL'),('External link imekataliwa','Tumia HTTPS URL halali'),('Network unit haitaki kuwa Published','Kagua fields na approval/consent declaration'),('AdSense unit imekataliwa','Kagua Publisher ID na Ad slot; usibandike script nzima'),('Adsterra unit imekataliwa','Kagua key, matching HTTPS URL na mwisho /KEY/invoke.js'),('Network ads hazionekani localhost','Hii ni tabia iliyokusudiwa; live ads hupimwa production'),('Script ready ipo bila tangazo','Kagua provider dashboard, approval, fill na browser blocking'),('Mapato hayaonekani admin','Reporting APIs bado hazijaunganishwa'),('Pop-up haijirudii','Kagua Pop-up frequency na browser storage')],[2.6,4.3])
p('Ukihitaji msaada, andika jina la placement/unit, public page iliyoathirika, muda wa tatizo kwa EAT na ujumbe wa error. Usitume password au siri za akaunti kwenye taarifa hiyo.')

page('11 Checklist kabla ya kuchapisha')
h('Direct sponsor')
check(['Jina na maudhui ya sponsor ni sahihi.','Picha inaonekana vizuri.','Destination link ni sahihi.','Kurasa zinazolengwa zimechaguliwa.','Position na priority zimekaguliwa.','Start na End ni sahihi kwa EAT.','Preview imekaguliwa.','Public page imekaguliwa baada ya kuchapisha.'])
h('Network ads')
check(['Public website imeandaliwa.','Provider approval inayohitajika imekamilika.','IDs, key na script URL zimetoka kwenye akaunti halisi.','Privacy na consent setup inayotakiwa imekamilika.','Main owner amekagua declaration.','Display pages, priority na schedule zimewekwa.','Live delivery imekaguliwa production.','Takwimu rasmi zinakaguliwa kwenye provider dashboard.'])

page('12 Daftari la sponsors')
p('Daftari hili ni sehemu ya document kwa usimamizi wa kazi. Si feature mpya ndani ya mfumo. Tumia rekodi moja kwa kila placement au network unit. Nakili ukurasa huu unapohitaji rekodi zaidi.')
for n in range(1,3):
    h(f'Rekodi {n}')
    table(['Taarifa','Nafasi ya kujaza'],[(x,'') for x in ['Sponsor / Provider','Placement / Unit name','Kurasa na nafasi','Priority na status','Start na End kwa EAT','Tarehe ya ukaguzi','Maelezo']],[2.4,4.5])

d.save(OUT)
print(OUT)
